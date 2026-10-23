#!/usr/bin/env python3
"""Export game/data/runes.json from design/runes-workbook.xlsx (design/14 section 7).

The workbook is Brent's and the source of truth for runes. This script only READS it, with
openpyxl data_only=True, so it gets the values Excel last computed and saved. It never writes to
the workbook. Edit the workbook (in Excel, and save it there so the cached values update), then
re-run:

    python tools/runes/export_runes_json.py

What it writes:
- base_runes: every row of "Base Runes" (id, affinity, name, character, opposite, role, valence,
  electronegativity, meaning = "As a skill"), plus "as_means" where the affinity sheet has one;
- universal_runes: every row of "Universal Runes" (meaning = "As a means"). Brent has not given
  them a role, valence or electronegativity yet, so they carry PLACEHOLDER values (below) and
  "placeholder": true;
- axes and affinity_positions: "Slider Links" columns G-K (-1..+1 per axis);
- rule_set / rule_sets: which rule set is in force (ACTIVE_RULE_SET) and each rule set's script and
  params. bond_v1's params are the alignment thresholds ("Alignment" B19/B20), the role table,
  electronegativity centre and bond-kind thresholds ("Bond Rules"), and the word rules from Brent's
  2026-10-22 rulings (WORD_RULES below - these are not in the workbook). bond_v2 (in force since
  2026-10-23, design/14 section 8) adds structure v2 from "Bond Rules" rows 43-59: the rings (fade
  and runes each seal adds, rows 47-51) and the shape traits (rows 55-59, their needs and changes
  read out of the cells' words), plus RING_RULES and REPEAT_RULES below. A new rule set gets an
  entry here and a script in game/scripts/runes/;
- skill_anchor: ability id (game/data/abilities.json) -> the base rune it is built on: each ability's
  "base_rune" (written by tools/skills/emit_abilities.py from the skill drafts, build brief
  2026-10-22). The affinity sheets' "Skills on it" names are still read, as a cross-check only:
  names with no ability, and sheet runes that disagree with an ability's base_rune, are listed under
  "unmatched" and printed.
"""
import json
import re
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "design/runes-workbook.xlsx"
ABILITIES = ROOT / "game/data/abilities.json"
OUT = ROOT / "game/data/runes.json"

# Universal runes have no role yet (Brent has not set them). Until he does they are Links at the
# electronegativity centre: two bonds, and an even bond with a middling rune. [MODEL placeholder]
UNIVERSAL_PLACEHOLDER = {"valence": 2, "electronegativity": 2.5}

# Brent 2026-10-22 (design/14 section 7 and the rune-setting brief). Not in the workbook.
WORD_RULES = OrderedDict([
    # level 0 holds 4 runes, level 1 holds 6: +2 per level [B]
    ("capacity_base", 4),
    ("capacity_per_level", 2),
    # OPEN [B 2026-10-22]: whether the skill's own base rune counts in the 4. Default: it does.
    ("anchor_counts", True),
    # not ruled [MODEL]: the same rune may appear more than once in a word (as atoms repeat in a
    # molecule); set false to forbid it
    ("allow_repeats", True),
    # not ruled [MODEL]: a closed word must also be one piece, every rune joined to the anchor
    ("closed_needs_connected", True),
])

# Which rule set the game uses. A rule set is a GDScript extending RuneRules; its entry under
# "rule_sets" names the script and carries its params. Interactions (RuneInteraction scripts in
# game/scripts/runes/interactions/) are listed by id per rule set.
ACTIVE_RULE_SET = "bond_v2"

# Structure v2 (Brent 2026-10-23, design/14 section 8; Bond Rules rows 43-59 hold the numbers). These
# are the parts the workbook does not hold. The rule set is game/scripts/runes/ring_rules.gd, its
# interaction game/scripts/runes/interactions/effects_v2.gd.
RING_RULES = OrderedDict([
    # OPEN [MODEL]: how "Runes it adds" (Bond Rules C47:C51) reads. "grows": each ring holds the ring
    # inside it plus its adds (3, 5, 7, 9, 11) - the only reading under which the great ring (a closed
    # ring of 5 or more, row 56) can ever form. "own": each ring holds exactly its adds (3, 2, 2, 2):
    # total +2 per level, as row 52 words it, but no ring past the first can be great.
    ("ring_capacity", "grows"),
    # [MODEL] the skill's own rune can only bond to ring 1, so it counts as filled once it holds as
    # many bonds as ring 1 has places (a four-bond rune would otherwise never close at level 1)
    ("anchor_fills_by_ring", True),
    # [MODEL] a seal is formed only from a closed word (the old lock rule, Bond Rules B36)
    ("seal_needs_closed", True),
])

# [MODEL] the generic "repeat once" a great ring gives: how long after the cast, and which kinds of
# skill cannot repeat (moving you, held, or lasting things) - for those the great ring's 40% lands
# as power through the fallback instead.
REPEAT_RULES = OrderedDict([
    ("delay_s", 0.45),
    ("no_repeat_kinds", ["dash", "blink", "grapple", "streak", "channel", "charge_release", "self_buff", "summon",
                         "rod", "prop", "seed"]),
])

BOND_V2_NOTES = [
    "Runes are sealed when a vestige (or a later construct) is bound into the skill; the word changes the skill only once sealed.",
    "Strength = the ring's fade x every bond's transmission on the strongest path to the skill.",
    "An ionic bond hands the weaker rune's whole effect to the stronger one.",
    "Universal runes have placeholder bonds (2) until Brent sets them.",
]

# Shown on the rune-setting screen under the word: what is not decided yet.
BOND_V1_NOTES = [
    "A locked word changes its skill; an unlocked one does nothing yet.",
    "Not decided yet: bond strength, unstable words, flaws, reactions between runes.",
    "Universal runes have placeholder bonds (2) until Brent sets them.",
]

# Rune effects, layers 1-3 (Brent 2026-10-22; constants on Bond Rules rows 19-40). These are the
# parts the workbook does not hold as numbers. The interaction script is
# game/scripts/runes/interactions/effects_v1.gd.
EFFECTS_V1 = OrderedDict([
    # Bond Rules A40: "a quarter of the change as power instead" - the quarter as a number
    ("fallback_fraction", 0.25),
    # [MODEL] a grant (a status, a shove, a mark...) the skill cannot use counts as this much change
    # before the quarter is taken: so an unusable grant at full strength is +5% power
    ("grant_as_change", 0.2),
    # [MODEL] where the fallback goes: power, or - for a skill that deals no damage - cost (a power
    # gain becomes a cost cut of the same size)
    ("fallback_groups", ["power", "cost"]),
    # groups where a smaller number is the better one (their fallback and their colour flip)
    ("lower_is_better", ["cost", "reflux", "windup"]),
    # stat group -> ability fields (abilities.json). "a.b" is a nested field; "field@kind|kind" only
    # counts for skills of those kinds. Built from what abilities.json really has (2026-10-22).
    ("stat_groups", OrderedDict([
        ("power", ["damage", "damage_mult", "on_hit.damage", "on_rise.damage", "relay_hit.damage", "zone.payload.damage",
                   "per_stack_damage", "bonus_damage", "pip_damage", "per_stack_burst", "damage_from_heat", "retaliate",
                   "on_pass.damage", "on_arrive.damage", "on_stop.damage", "on_hit.consume.per_stack_damage", "per_target_damage",
                   "ray_damage", "corpse_burst.damage", "on_break.damage", "damage_full", "damage_min", "damage_mult_min",
                   "zone.on_enter.damage", "thrall.damage"]),
        ("reach", ["range", "life@projectile", "reach_bonus", "range_bonus", "radius", "burst_radius", "arc_deg", "width@beam|prop",
                   "jump_radius", "pass_radius", "relay_radius", "relay_range", "arrive_radius", "through_dist",
                   "zone.radius", "zone.length", "zone.width", "trail.width", "leave_zone.radius", "light_radius", "gather_radius",
                   "reveal_radius", "leave_decoy.lure_radius", "corpse_range", "drag_radius", "trigger_radius"]),
        ("speed", ["speed", "dash_speed"]),
        ("cost", ["cost"]),
        ("reflux", ["reflux_per_cost"]),
        ("windup", ["windup", "recovery", "delay", "rise_time"]),
        ("duration", ["duration", "life@prop|rod", "zone.life", "trail.life", "leave_prop.life", "on_hit.status_s", "on_pass.status_s",
                      "leave_zone.life", "leave_decoy.life", "burst_zone.life", "thrall.life"]),
        # 2026-10-22: poise on a rise or an arrival counts too
        ("poise", ["on_hit.poise", "heavy_poise", "on_stop.poise", "zone.payload.poise", "on_rise.poise", "on_arrive.poise",
                   "on_pass.poise", "per_stack_poise"]),
    ])),
    # where a grant goes: the first of these hit payloads the skill has
    ("payload_parts", ["on_hit", "on_arrive", "on_pass", "on_stop", "on_rise", "relay_hit", "zone.payload", "trail.payload"]),
    # a field that is a full turn can grow no further
    ("field_caps", OrderedDict([("arc_deg", 360)])),
    # the skill's stage multiplies these when it is cast (Player.ability_damage_mult), so the preview does too
    ("stage_scaled_groups", ["power"]),
    # units for the effects preview
    ("group_units", OrderedDict([("windup", " s"), ("duration", " s")])),
    ("field_units", OrderedDict([("arc_deg", " deg"), ("life", " s"), ("range", " m"), ("radius", " m"), ("reach_bonus", " m"),
                                 ("range_bonus", " m"), ("burst_radius", " m"), ("jump_radius", " m"), ("zone.radius", " m"),
                                 ("zone.length", " m"), ("zone.width", " m"), ("width", " m")])),
    # words for the effects preview
    ("group_words", OrderedDict([("power", "damage"), ("reach", "range and area"), ("speed", "speed"), ("cost", "essence cost"),
                                 ("reflux", "reflux"), ("windup", "wind-up"), ("duration", "duration"), ("poise", "poise damage")])),
    ("field_labels", OrderedDict([
        ("damage", "Damage"), ("damage_mult", "Damage (x a strike)"), ("on_hit.damage", "Damage"), ("on_rise.damage", "Damage as it rises"),
        ("relay_hit.damage", "Relay damage"), ("zone.payload.damage", "Damage per tick"), ("per_stack_damage", "Damage per stack"),
        ("bonus_damage", "Bonus damage"), ("pip_damage", "Damage per pip"), ("per_stack_burst", "Burst per stack"),
        ("damage_from_heat", "Damage per heat"), ("retaliate", "Retaliation damage"),
        ("range", "Range"), ("life@projectile", "Flight time"), ("reach_bonus", "Reach"), ("range_bonus", "Reach"), ("radius", "Radius"),
        ("burst_radius", "Burst radius"), ("arc_deg", "Arc"), ("width@beam|prop", "Width"), ("jump_radius", "Jump radius"),
        ("pass_radius", "Sweep radius"), ("relay_radius", "Relay radius"), ("relay_range", "Relay range"), ("arrive_radius", "Arrival radius"),
        ("through_dist", "Carries through"), ("zone.radius", "Zone radius"), ("zone.length", "Zone length"), ("zone.width", "Zone width"),
        ("trail.width", "Trail width"), ("speed", "Speed"), ("dash_speed", "Dash speed"), ("cost", "Essence cost"),
        ("reflux_per_cost", "Reflux per essence"), ("windup", "Wind-up"), ("recovery", "Recovery"), ("delay", "Delay"),
        ("rise_time", "Rise time"), ("duration", "Duration"), ("life@prop|rod", "Stands for"), ("zone.life", "Zone lasts"),
        ("trail.life", "Trail lasts"), ("leave_prop.life", "Stone stands for"), ("on_hit.poise", "Poise damage"),
        ("heavy_poise", "Poise damage (heavy)"), ("on_stop.poise", "Poise damage"), ("zone.payload.poise", "Poise damage per tick"),
    ])),
])

EFFECT_GROUPS = list(EFFECTS_V1["stat_groups"].keys())
STATUSES = ROOT / "game/data/statuses.json"


def parse_effect(text_in, rid, statuses):
    """One op per ';' (Bond Rules A40). Returns a list of ops; exits on anything it cannot read."""
    ops = []
    for part in [p.strip() for p in str(text_in or "").split(";") if p.strip()]:
        m = re.fullmatch(r"(\w+) x([\d.]+)", part)
        if m and m.group(1) in EFFECT_GROUPS:
            ops.append(OrderedDict([("op", "mul"), ("group", m.group(1)), ("value", clean(float(m.group(2))))]))
            continue
        m = re.fullmatch(r"status (\w+)(?: (\d+))?", part)
        if m:
            if m.group(1) not in statuses:
                sys.exit("export_runes_json: %s grants status '%s', which statuses.json does not have" % (rid, m.group(1)))
            ops.append(OrderedDict([("op", "status"), ("id", m.group(1)), ("stacks", int(m.group(2) or 1))]))
            continue
        m = re.fullmatch(r"(shove|pull|marks|jumps|pierce) \+([\d.]+)", part)
        if m:
            ops.append(OrderedDict([("op", m.group(1)), ("value", clean(float(m.group(2))))]))
            continue
        # through (2026-10-22): the hit ignores front block and shields
        if part in ("interrupts", "braced", "through"):
            ops.append(OrderedDict([("op", part)]))
            continue
        m = re.fullmatch(r"bonus (\w+) \+([\d.]+)", part)
        if m:
            if m.group(1) not in statuses:
                sys.exit("export_runes_json: %s has a bonus against '%s', which statuses.json does not have" % (rid, m.group(1)))
            ops.append(OrderedDict([("op", "bonus"), ("status", m.group(1)), ("value", clean(float(m.group(2))))]))
            continue
        sys.exit("export_runes_json: %s effect '%s' is not in the effect syntax (Bond Rules A40)" % (rid, part))
    return ops


def text(v):
    return "" if v is None else str(v).strip()


def read_structure_v2(cells):
    """Bond Rules rows 43-59 (Brent 2026-10-23): the rings (fade, runes each seal adds) and the shape
    traits. The traits' needs and changes are read out of their words; anything that does not read
    stops the export, so a reworded cell is noticed rather than silently dropped."""
    if not text(cells.get((43, 1))).lower().startswith("structure v2"):
        sys.exit("export_runes_json: Bond Rules A43 is not the structure v2 header - the layout moved")
    if text(cells.get((46, 1))).lower() != "ring (seal)" or text(cells.get((46, 2))).lower() != "fade":
        sys.exit("export_runes_json: Bond Rules A46:B46 are not 'Ring (seal)', 'Fade' - the layout moved")
    rings = []
    for r in range(47, 52):
        label = text(cells.get((r, 1)))
        m = re.match(r"(\d+)", label)
        if not m:
            sys.exit("export_runes_json: Bond Rules A%d '%s' does not start with a ring number" % (r, label))
        rings.append(OrderedDict([("ring", int(m.group(1))), ("or_more", "more" in label.lower()),
                                  ("fade", clean(num(cells.get((r, 2)), "Bond Rules B%d" % r))),
                                  ("adds", int(num(cells.get((r, 3)), "Bond Rules C%d" % r)))]))
    if [x["ring"] for x in rings] != [1, 2, 3, 4, 5] or not rings[-1]["or_more"]:
        sys.exit("export_runes_json: Bond Rules A47:A51 are not rings 1, 2, 3, 4, 5 or more - the layout moved")
    if text(cells.get((54, 1))).lower() != "shape trait":
        sys.exit("export_runes_json: Bond Rules A54 is not 'Shape trait' - the layout moved")
    rows = {}
    for r in range(55, 60):
        rows[text(cells.get((r, 1))).lower()] = (r, text(cells.get((r, 2))), text(cells.get((r, 3))), text(cells.get((r, 4))))
    want = ["closed ring", "great ring", "spoke", "branch", "balanced"]
    if list(rows) != want:
        sys.exit("export_runes_json: Bond Rules A55:A59 are not %s - the layout moved" % want)

    def must(pattern, value, where):
        m = re.search(pattern, value)
        if not m:
            sys.exit("export_runes_json: Bond Rules %s '%s' does not read as /%s/" % (where, value, pattern))
        return m

    def group(word):
        g = word.lower()
        if g not in EFFECT_GROUPS:
            sys.exit("export_runes_json: shape trait group '%s' is not a stat group %s" % (word, EFFECT_GROUPS))
        return g

    traits = OrderedDict()
    r, needs, change, what = rows["closed ring"]
    m = must(r"(\w+) x([\d.]+) per closed ring", change, "C%d" % r)
    traits["closed_ring"] = OrderedDict([("label", "Closed ring"), ("needs", needs), ("change", change), ("what", what),
                                         ("min_runes", 3), ("ops", [OrderedDict([("op", "mul"), ("group", group(m.group(1))), ("value", clean(float(m.group(2))))])])])
    r, needs, change, what = rows["great ring"]
    m1 = must(r"closed ring of (\d+) or more", needs, "B%d" % r)
    m2 = must(r"repeats (once|twice) at (\d+)%", change, "C%d" % r)
    traits["great_ring"] = OrderedDict([("label", "Great ring"), ("needs", needs), ("change", change), ("what", what),
                                        ("min_runes", int(m1.group(1))), ("repeats", 1 if m2.group(1) == "once" else 2),
                                        ("mult", clean(int(m2.group(2)) / 100.0))])
    r, needs, change, what = rows["spoke"]
    m1 = must(r"through (\d+) or more rings", needs, "B%d" % r)
    m2 = must(r"(\w+) x([\d.]+) and (\w+) x([\d.]+) per ring past the (first|second|third)", change, "C%d" % r)
    traits["spoke"] = OrderedDict([("label", "Spoke"), ("needs", needs), ("change", change), ("what", what),
                                   ("min_rings", int(m1.group(1))), ("free_rings", {"first": 1, "second": 2, "third": 3}[m2.group(5)]),
                                   ("ops", [OrderedDict([("op", "mul"), ("group", group(m2.group(1))), ("value", clean(float(m2.group(2))))]),
                                            OrderedDict([("op", "mul"), ("group", group(m2.group(3))), ("value", clean(float(m2.group(4))))])])])
    r, needs, change, what = rows["branch"]
    m1 = must(r"(\d+) or more bonds", needs, "B%d" % r)
    m2 = must(r"(\w+)(?: \(\w+\))? x([\d.]+) per branch rune, up to (\d+)", change, "C%d" % r)
    traits["branch"] = OrderedDict([("label", "Branch"), ("needs", needs), ("change", change), ("what", what),
                                    ("min_bonds", int(m1.group(1))), ("max_count", int(m2.group(3))),
                                    ("ops", [OrderedDict([("op", "mul"), ("group", group(m2.group(1))), ("value", clean(float(m2.group(2))))])])])
    r, needs, change, what = rows["balanced"]
    must(r"ring 1 full", needs, "B%d" % r)
    must(r"same kind", needs, "B%d" % r)
    must(r"keeps all", change, "C%d" % r)
    traits["balanced"] = OrderedDict([("label", "Balanced"), ("needs", needs), ("change", change), ("what", what), ("anchor_factor", 1)])
    return rings, traits, " ".join([text(cells.get((44, 1))), text(cells.get((52, 1)))])


def num(v, what):
    if v is None or v == "":
        sys.exit("export_runes_json: %s is empty in the workbook (open it in Excel and save, so values are cached)" % what)
    return float(v)


def clean(x):
    """Round away float noise from Excel (2.6199999999 -> 2.62); keep ints as ints."""
    x = round(float(x), 4)
    return int(x) if x == int(x) else x


def main():
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)

    # --- bond rules ---------------------------------------------------------------------
    br = wb["Bond Rules"]
    cells = {(r, c): v for r, row in enumerate(br.iter_rows(min_row=1, max_row=60, values_only=True), 1) for c, v in enumerate(row, 1)}
    roles = OrderedDict()
    for r in range(4, 8):
        role = text(cells.get((r, 1)))
        roles[role.lower()] = OrderedDict([("valence", int(num(cells.get((r, 2)), "Bond Rules B%d" % r))),
                                           ("spread", clean(num(cells.get((r, 3)), "Bond Rules C%d" % r)))])
    if text(cells.get((9, 1))).lower() != "electronegativity centre":
        sys.exit("export_runes_json: Bond Rules A9 is not 'Electronegativity centre' - the layout moved")
    en_centre = clean(num(cells.get((9, 2)), "Bond Rules B9"))
    bond_kinds = []
    for r in range(13, 16):
        bond_kinds.append(OrderedDict([("kind", text(cells.get((r, 1))).lower()), ("from", clean(num(cells.get((r, 2)), "Bond Rules B%d" % r)))]))
    if [k["kind"] for k in bond_kinds] != ["even", "polar", "ionic"]:
        sys.exit("export_runes_json: Bond Rules A13:A15 are not Even, Polar, Ionic - the layout moved")

    # --- rune effects, layers 1-3 (Bond Rules rows 19-40) ---------------------------------------
    if not text(cells.get((22, 1))).lower().startswith("bonds from the skill"):
        sys.exit("export_runes_json: Bond Rules A22 is not the strength-by-bonds header - the layout moved")
    fade = []
    for r in range(23, 27):
        bonds = text(cells.get((r, 1)))
        fade.append(OrderedDict([("bonds", int(re.match(r"\d+", bonds).group(0))), ("or_more", "more" in bonds),
                                 ("strength", clean(num(cells.get((r, 2)), "Bond Rules B%d" % r)))]))
    kind_factors = OrderedDict()
    for r in range(29, 32):
        kind_factors[text(cells.get((r, 1))).lower()] = OrderedDict([("more_pull", clean(num(cells.get((r, 2)), "Bond Rules B%d" % r))),
                                                                     ("less_pull", clean(num(cells.get((r, 3)), "Bond Rules C%d" % r)))])
    if list(kind_factors) != ["even", "polar", "ionic"]:
        sys.exit("export_runes_json: Bond Rules A29:A31 are not Even, Polar, Ionic - the layout moved")
    labels = {33: "a status applies", 34: "lowest multiplier", 35: "highest multiplier", 36: "a word must be closed"}
    for r, start in labels.items():
        if not text(cells.get((r, 1))).lower().startswith(start):
            sys.exit("export_runes_json: Bond Rules A%d is not '%s...' - the layout moved" % (r, start))
    status_threshold = clean(num(cells.get((33, 2)), "Bond Rules B33"))
    mult_min = clean(num(cells.get((34, 2)), "Bond Rules B34"))
    mult_max = clean(num(cells.get((35, 2)), "Bond Rules B35"))
    lock_needs_closed = bool(num(cells.get((36, 2)), "Bond Rules B36"))
    anchor_rule_text = text(cells.get((38, 1)))
    syntax_text = text(cells.get((40, 1)))
    statuses = set(json.loads(STATUSES.read_text(encoding="utf-8"))["statuses"].keys())
    rings, shape_traits, structure_text = read_structure_v2(cells)

    # --- alignment thresholds -------------------------------------------------------------
    al = wb["Alignment"]
    al_cells = {(r, c): v for r, row in enumerate(al.iter_rows(min_row=1, max_row=22, max_col=2, values_only=True), 1) for c, v in enumerate(row, 1)}
    if not text(al_cells.get((19, 1))).lower().startswith("refuse") or not text(al_cells.get((20, 1))).lower().startswith("weaken"):
        sys.exit("export_runes_json: Alignment A19/A20 are not the refuse / weaken thresholds - the layout moved")
    refuse_below = clean(num(al_cells.get((19, 2)), "Alignment B19"))
    weaken_below = clean(num(al_cells.get((20, 2)), "Alignment B20"))

    # --- affinity positions -----------------------------------------------------------------
    sl = wb["Slider Links"]
    rows = list(sl.iter_rows(min_row=1, max_row=17, values_only=True))
    axes = []
    for h in rows[0][6:11]:
        axes.append(re.sub(r"\s*\(.*\)", "", text(h)).lower().replace("/", "_").replace(" ", "_"))
    positions = OrderedDict()
    for row in rows[1:]:
        if text(row[0]):
            positions[text(row[0]).lower()] = [clean(num(v, "Slider Links %s" % text(row[0]))) for v in row[6:11]]

    # --- base runes ---------------------------------------------------------------------------
    base = OrderedDict()
    for row in wb["Base Runes"].iter_rows(min_row=2, values_only=True):
        rid = text(row[0])
        if not rid:
            continue
        role = text(row[10]).lower()
        if role not in roles:
            sys.exit("export_runes_json: %s has role '%s', not one of %s" % (rid, row[10], list(roles)))
        valence = int(num(row[12], "%s valence" % rid))
        if valence != roles[role]["valence"]:
            sys.exit("export_runes_json: %s valence %d does not match its role %s" % (rid, valence, role))
        base[rid] = OrderedDict([
            ("id", rid), ("affinity", text(row[1]).lower()), ("name", text(row[3])), ("character", text(row[4])),
            ("opposite", text(row[5])), ("role", role), ("valence", valence),
            ("electronegativity", clean(num(row[13], "%s electronegativity" % rid))), ("meaning", text(row[8])),
        ])
        effect = text(row[14]) if len(row) > 14 else ""
        if effect:
            base[rid]["effect"] = effect
            base[rid]["ops"] = parse_effect(effect, rid, statuses)

    # --- universal runes ------------------------------------------------------------------------
    universal = OrderedDict()
    for row in wb["Universal Runes"].iter_rows(min_row=2, values_only=True):
        rid = text(row[0])
        if not rid:
            continue
        universal[rid] = OrderedDict([
            ("id", rid), ("name", text(row[1])), ("character", text(row[2])), ("opposite", text(row[3])),
            ("valence", UNIVERSAL_PLACEHOLDER["valence"]), ("electronegativity", UNIVERSAL_PLACEHOLDER["electronegativity"]),
            ("placeholder", True), ("meaning", text(row[6])),
        ])
        effect = text(row[7]) if len(row) > 7 else ""
        if effect:
            universal[rid]["effect"] = effect
            universal[rid]["ops"] = parse_effect(effect, rid, statuses)

    # --- affinity sheets: as-means text and the skills on each rune ----------------------------------
    abilities = json.loads(ABILITIES.read_text(encoding="utf-8"))["abilities"]
    by_name = {}
    for aid, a in abilities.items():
        by_name.setdefault(a["name"], []).append(aid)
    anchor = OrderedDict()
    unmatched_names = []      # "Skills on it" names with no ability of that name and affinity
    conflicts = []
    for aff in positions:
        sheet = aff.capitalize()
        if sheet not in wb.sheetnames:
            continue
        for row in wb[sheet].iter_rows(min_row=6, values_only=True):
            rid = text(row[1])
            if rid not in base:
                continue
            if text(row[5]):
                base[rid]["as_means"] = text(row[5])
            for raw in text(row[8]).split(","):
                name = re.sub(r"\s*\(.*?\)", "", raw).strip()
                if not name:
                    continue
                ids = [i for i in by_name.get(name, []) if abilities[i].get("affinity") == base[rid]["affinity"]]
                if len(ids) != 1:
                    unmatched_names.append("%s (%s, %s)" % (name, rid, base[rid]["name"]))
                    continue
                if ids[0] in anchor and anchor[ids[0]] != rid:
                    conflicts.append("%s on both %s and %s" % (ids[0], anchor[ids[0]], rid))
                    continue
                anchor[ids[0]] = rid
    # the anchor is the ability's own base_rune (build brief 2026-10-22); the sheet is only a cross-check
    sheet = anchor
    anchor = OrderedDict()
    for aid, a in abilities.items():
        rid = str(a.get("base_rune", ""))
        if rid in base:
            anchor[aid] = rid
            if aid in sheet and sheet[aid] != rid:
                conflicts.append("%s: base_rune %s, but the %s sheet puts it on %s" % (aid, rid, a.get("affinity", "?").capitalize(), sheet[aid]))
        elif aid in sheet:
            anchor[aid] = sheet[aid]
    no_rune = [aid for aid in abilities if aid not in anchor]

    out = OrderedDict([
        ("_note", "GENERATED by tools/runes/export_runes_json.py from design/runes-workbook.xlsx - do not hand-edit. "
                  "Edit the workbook (and save it in Excel so its values are cached), or WORD_RULES in the script, and re-run. "
                  "Rune words: design/14 section 7 (Brent 2026-10-22). A skill is built on one base rune (skill_anchor); "
                  "runes are added and linked. A saved word is rule-free (nodes + links, scripts/runes/rune_word.gd); "
                  "'rule_set' picks the rules in force from 'rule_sets' (each names a script extending RuneRules, plus its params). "
                  "bond_v1 (scripts/runes/bond_rules.gd): single bonds; each rune makes at most its valence in bonds; a word is closed "
                  "when every rune's bonds equal its valence. The electronegativity difference sets a bond's kind (bond_kinds, "
                  "'from' = smallest difference). Two base runes of different affinities bond only if the cosine of their "
                  "affinity_positions is at least refuse_below; below weaken_below the bond is weakened. Universal "
                  "runes have no position. Universal valence and electronegativity are PLACEHOLDERS (placeholder: true): Brent "
                  "has not set them. Rune effects (layers 1-3, Bond Rules rows 19-40): each rune's 'effect' text and parsed 'ops', worked out by "
                  "the interaction effects_v1 (scripts/runes/interactions/effects_v1.gd) from bond_v1.interaction_params.effects_v1; "
                  "they change a skill only while its word is locked. "
                  "bond_v2 (scripts/runes/ring_rules.gd, Brent 2026-10-23, design/14 section 8, Bond Rules rows 43-59) is in force: "
                  "the word is concentric rings round the skill, one per level, sealed when a construct is bound into the skill; "
                  "'rings' gives each ring's fade and the runes it adds (ring_capacity says how 'adds' reads); a rune bonds only on "
                  "its own ring and the rings next to it; strength = the ring's fade x the product of the bonds' transmissions on "
                  "the strongest path; an ionic bond hands the weaker rune's effect to the stronger; shape_traits (closed ring, "
                  "great ring, spoke, branch, balanced) are worked out by effects_v2 (scripts/runes/interactions/effects_v2.gd)."),
        ("rule_set", ACTIVE_RULE_SET),
        ("rule_sets", OrderedDict([
            ("bond_v1", OrderedDict([
                ("script", "res://scripts/runes/bond_rules.gd"),
                ("title", "Bonds (first rules, 2026-10-22)"),
                ("refuse_below", refuse_below), ("weaken_below", weaken_below),
                ("electronegativity_centre", en_centre), ("roles", roles), ("bond_kinds", bond_kinds),
            ] + list(WORD_RULES.items()) + [
                ("lock_needs_closed", lock_needs_closed),
                ("interactions", ["effects_v1"]),
                ("interaction_params", OrderedDict([
                    ("effects_v1", OrderedDict([
                        ("fade", fade), ("kind_factors", kind_factors), ("status_threshold", status_threshold),
                        ("mult_min", mult_min), ("mult_max", mult_max),
                    ] + list(EFFECTS_V1.items()) + [
                        ("anchor_rule", anchor_rule_text), ("syntax", syntax_text),
                    ])),
                ])),
                ("notes", BOND_V1_NOTES),
            ])),
            ("bond_v2", OrderedDict([
                ("script", "res://scripts/runes/ring_rules.gd"),
                ("title", "Seals (structure v2, 2026-10-23)"),
                ("refuse_below", refuse_below), ("weaken_below", weaken_below),
                ("electronegativity_centre", en_centre), ("roles", roles), ("bond_kinds", bond_kinds),
            ] + list(WORD_RULES.items()) + [
                ("lock_needs_closed", lock_needs_closed),
                ("rings", rings),
            ] + list(RING_RULES.items()) + [
                ("structure", structure_text),
                ("interactions", ["effects_v2"]),
                ("interaction_params", OrderedDict([
                    ("effects_v2", OrderedDict([
                        ("fade", fade), ("kind_factors", kind_factors), ("status_threshold", status_threshold),
                        ("mult_min", mult_min), ("mult_max", mult_max),
                    ] + list(EFFECTS_V1.items()) + [
                        ("ring_fade", [OrderedDict([("ring", r["ring"]), ("or_more", r["or_more"]), ("fade", r["fade"])]) for r in rings]),
                        ("shape_traits", shape_traits),
                        ("repeat", REPEAT_RULES),
                        ("anchor_rule", anchor_rule_text), ("syntax", syntax_text),
                    ])),
                ])),
                ("notes", BOND_V2_NOTES),
            ])),
        ])),
        ("axes", axes),
        ("affinity_positions", positions),
        ("base_runes", base),
        ("universal_runes", universal),
        ("skill_anchor", anchor),
        ("unmatched", OrderedDict([("sheet_skills_without_ability", unmatched_names), ("abilities_without_rune", no_rune),
                                   ("conflicts", conflicts)])),
    ])
    OUT.write_text(json.dumps(out, indent="\t", ensure_ascii=False) + "\n", encoding="utf-8", newline="\n")
    print("wrote %s: %d base runes, %d universal runes, %d affinities, %d skills anchored" % (
        OUT.relative_to(ROOT), len(base), len(universal), len(positions), len(anchor)))
    for n in unmatched_names:
        print("  unmatched sheet skill: " + n)
    for a in no_rune:
        print("  ability with no rune: " + a)
    for c in conflicts:
        print("  conflict: " + c)


if __name__ == "__main__":
    main()
