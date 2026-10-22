"""Build design/runes-workbook.xlsx from design/14 (first passes) and design/affinities/water.md (deep dive)."""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = 'E:/Code/lifesteal-crawler/'
OUT = ROOT + 'design/runes-workbook.xlsx'

ORDER = [  # name, code, tab colour
    ('Water', 'WAT', '3A7BD5'), ('Fire', 'FIR', 'D9534F'), ('Earth', 'EAR', '8B6B3E'), ('Wind', 'WIN', '7FB8A4'),
    ('Lightning', 'LIG', 'E8C547'), ('Ice', 'ICE', 'A8D8EA'), ('Dark', 'DRK', '3B3355'), ('Metal', 'MET', '8C9399'),
    ('Blood', 'BLD', '8B1E2D'), ('Life', 'LIF', '6BBF59'), ('Death', 'DTH', '4A4A4A'), ('Nature', 'NAT', '3E7B3A'),
    ('Light', 'LGT', 'F5E6A8'), ('Dream', 'DRM', '9B7FD4'), ('Poison', 'POI', '7CB342'), ('Disease', 'DIS', '6B6B3A'),
]
PROPOSED = ['Sound', 'Gravity', 'Space', 'Time', 'Mind', 'Beast', 'Spirit']

# --- parse design/14 -------------------------------------------------------------------------
d14 = open(ROOT + 'design/14-rune-sets-v0.md', encoding='utf-8').read()
CELL = re.compile(r'\*\*(.+?)\*\*:\s*(.+)')


def split_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


affinities = {}
for m in re.finditer(r'^### (\w+) — (.+?)$', d14, re.M):
    name = m.group(1)
    head = m.group(2)
    body = d14[m.end():]
    nxt = re.search(r'^(###|## |\*\*Totals)', body, re.M)
    body = body[:nxt.start()] if nxt else body
    prose = ' '.join(l.strip() for l in body.splitlines() if l.strip() and not l.strip().startswith('|')).replace('**', '')
    rule = re.sub(r'\s*\[[^\]]*\]', '', head).strip()
    source = '; '.join(re.findall(r'\[([^\]]*)\]', head))
    pairs = []
    for line in body.splitlines():
        if not line.startswith('| **'):
            continue
        cells = split_row(line)
        a, b = CELL.match(cells[0]), CELL.match(cells[1])
        skills = cells[2] if len(cells) > 2 else ''
        sa, sb = (skills.split(' / ') + [''])[:2] if skills else ('', '')
        clean = lambda x: '' if x.strip() in ('—', '-') else x.strip()
        pairs.append([(a.group(1), a.group(2), clean(sa)), (b.group(1), b.group(2), clean(sb))])
    affinities[name] = {'rule': rule[0].upper() + rule[1:], 'note': prose, 'source': source, 'pairs': pairs,
                        'status': 'First pass', 'means': {}, 'together': {}}

# --- water from its deep dive ----------------------------------------------------------------
wd = open(ROOT + 'design/affinities/water.md', encoding='utf-8').read()
wpairs, means, together = [], {}, {}
sec = wd[wd.index('## 4. Base runes'):wd.index('## 5. Skills')]
for line in sec.splitlines():
    if not line.startswith('| **'):
        continue
    c = split_row(line)
    ra, rb = c[0].strip('*'), c[3].strip('*')
    tog = c[6]
    wpairs.append([(ra, c[1], ''), (rb, c[4], '')])
    means[ra], means[rb] = c[2], c[5]
    together[ra] = together[rb] = tog
water_skills = {'Surge': 'Water Bolt (starter), Breaker (starter)', 'Ebb': 'Undertow, Riptide Lash',
                'Flow': 'Sluice Step (starter)', 'Pool': 'Standing Pool', 'Crush': 'Pressure Jet', 'Coat': 'Backwash'}
water_new = {'Well': 'Proposed skill: Draw Water', 'Wring': 'Proposed skill: Wringing Grip',
             'Mist': 'Proposed skill: Fogbank', 'Mirror': 'Proposed skill: Glass Water',
             'Drown': 'Proposed skill: Drowning Hold'}
affinities['Water'] = {
    'rule': 'Moves bodies (shove and Crush); works the water that is already there (wetness)',
    'note': 'Deep dive: design/affinities/water.md', 'source': 'existing (design/05); MODEL: wetness, from C',
    'pairs': [[(n, s, water_skills.get(n, '')) for (n, s, _) in p] for p in wpairs],
    'status': 'Deep dive v0', 'means': means, 'together': together, 'new': water_new}
for k in ('Ice', 'Dark'):
    affinities[k]['status'] = 'First pass (new 2026-10-22)'
assert [a for a, _, _ in ORDER] == [a for a, _, _ in ORDER if a in affinities], set(a for a, _, _ in ORDER) - set(affinities)

# --- universal runes from design/14 §2 ------------------------------------------------------
usec = d14[d14.index('## 2. Universal runes'):d14.index('## 3. The affinities')]
universal = []
for line in usec.splitlines():
    if line.startswith('| **'):
        c = split_row(line)
        universal.append(((c[0].strip('*'), c[1]), (c[2].strip('*'), c[3])))
UNI_CHARS = list('⇢⊡»«↷⇥↑↓◯◉⊚⋯⊙⁘◆⁖▣□⌖⤳∞◌⧗⇡↻⊣⟲⟳⧓⫽⇇⇉⊳⊲☉⚬⊞⁂◇◈☼◒≈≋⇑⇓⊛⊘✦✧⚘⚱✵⬢⊜⚄')
assert len(universal) == 28 and len(UNI_CHARS) == 56

# --- characters for base runes: runic block for the five built affinities, alchemical for the rest
runic = [chr(c) for c in range(0x16A0, 0x16EB)]
alch = [chr(c) for c in range(0x1F700, 0x1F774)]
BUILT = {'Water', 'Fire', 'Earth', 'Wind', 'Lightning'}

# --- styles ---------------------------------------------------------------------------------
F = 'Arial'
f_body = Font(name=F, size=10)
f_input = Font(name=F, size=10, color='0000FF')
f_link = Font(name=F, size=10, color='008000')
f_calc = Font(name=F, size=10, color='000000')
f_head = Font(name=F, size=10, bold=True, color='FFFFFF')
f_title = Font(name=F, size=16, bold=True)
f_sub = Font(name=F, size=10, italic=True, color='555555')
fill_head = PatternFill('solid', fgColor='2F2A24')
fill_todo = PatternFill('solid', fgColor='FFFF00')
thin = Side(style='thin', color='D0D0D0')
box = Border(bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='center')


def char_font(ch, color='0000FF', size=18):
    name = 'Segoe UI Historic' if ch and 0x16A0 <= ord(ch[0]) <= 0x16FF else 'Segoe UI Symbol'
    return Font(name=name, size=size, color=color)


def header(ws, row, cols, widths):
    for i, (h, w) in enumerate(zip(cols, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.alignment = f_head, fill_head, Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 30


wb = Workbook()
readme = wb.active
readme.title = 'Read Me'
index = wb.create_sheet('Affinities')
master = wb.create_sheet('Base Runes')
uni = wb.create_sheet('Universal Runes')

# --- Base Runes (master) ----------------------------------------------------------------------
MCOLS = ['ID', 'Affinity', 'Pair', 'Rune', 'Character', 'Opposite ID', 'Opposite', 'Opposite character',
         'As a skill (from the affinity sheet)', 'Code point as first assigned']
header(master, 1, MCOLS, [10, 12, 6, 16, 11, 11, 16, 11, 60, 14])
master.freeze_panes = 'B2'
rows = []          # (id, affinity, pair, rune, char, opp_id, skill, means, together, skills, notes)
ri = ai = 0
for aff, code, _ in ORDER:
    a = affinities[aff]
    for p, pair in enumerate(a['pairs'], start=1):
        for side, (rune, skill, skills) in zip('ab', pair):
            if aff in BUILT:
                ch = runic[ri]; ri += 1
            else:
                ch = alch[ai]; ai += 1
            rid, oid = f'{code}-{p}{side}', f'{code}-{p}{"b" if side == "a" else "a"}'
            note = a.get('new', {}).get(rune, '')
            rows.append((rid, aff, p, rune, ch, oid, skill, a['means'].get(rune, ''), a['together'].get(rune, ''),
                         skills, note))
ids = [r[0] for r in rows]
assert len(ids) == len(set(ids)) == 166, len(ids)
names = [r[3] for r in rows]
dupes = sorted({n for n in names if names.count(n) > 1})

# where each rune's "as a skill" lives on its affinity sheet (filled while writing those sheets)
skill_cell = {}

# --- affinity sheets --------------------------------------------------------------------------
ACOLS = ['Pair', 'ID', 'Rune', 'Character', 'As a skill (the verb a skill is built on)',
         'As a means (added to another skill)', 'Opposite', 'Together (rune + opposite release)',
         'Skills on it', 'Notes']
for aff, code, colour in ORDER:
    a = affinities[aff]
    ws = wb.create_sheet(aff)
    ws.sheet_properties.tabColor = colour
    ws['A1'] = aff
    ws['A1'].font = f_title
    ws['A2'] = 'Rule: ' + a['rule']
    ws['A2'].font = Font(name=F, size=11, bold=True)
    ws['A3'] = (a['note'] + '  ' if a['note'] else '') + f"Status: {a['status']}." + (f"  Source: {a['source']}." if a['source'] else '')
    ws['A3'].font = f_sub
    for r in (2, 3):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[3].height = 30
    header(ws, 5, ACOLS, [6, 10, 16, 11, 46, 46, 16, 40, 30, 30])
    ws.freeze_panes = 'C6'
    r = 6
    for row in [x for x in rows if x[1] == aff]:
        rid, _, p, rune, ch, oid, skill, mean, tog, skills, note = row
        ws.cell(row=r, column=1, value=p).font = f_body
        ws.cell(row=r, column=2, value=rid).font = f_input
        c = ws.cell(row=r, column=3, value=f"=IFERROR(INDEX('Base Runes'!$D:$D,MATCH($B{r},'Base Runes'!$A:$A,0)),\"\")")
        c.font = Font(name=F, size=10, bold=True, color='008000')
        c = ws.cell(row=r, column=4, value=f"=IFERROR(INDEX('Base Runes'!$E:$E,MATCH($B{r},'Base Runes'!$A:$A,0)),\"\")")
        c.font, c.alignment = char_font(ch, '008000'), center
        ws.cell(row=r, column=5, value=skill).font = f_input
        c = ws.cell(row=r, column=6, value=mean or None)
        c.font = f_input
        if not mean:
            c.fill = fill_todo
        ws.cell(row=r, column=7, value=f"=IFERROR(INDEX('Base Runes'!$G:$G,MATCH($B{r},'Base Runes'!$A:$A,0)),\"\")").font = f_link
        c = ws.cell(row=r, column=8, value=tog or None)
        c.font = f_input
        if not tog:
            c.fill = fill_todo
        ws.cell(row=r, column=9, value=skills or None).font = f_input
        ws.cell(row=r, column=10, value=note or None).font = f_input
        for col in range(1, 11):
            ws.cell(row=r, column=col).alignment = center if col == 4 else wrap
            ws.cell(row=r, column=col).border = box
        ws.row_dimensions[r].height = 30
        skill_cell[rid] = f"='{aff}'!$E${r}" if ' ' in aff else f"={aff}!$E${r}"
        r += 1
    ws.cell(row=r + 1, column=1, value='Blue: type here.  Green: pulled from Base Runes (rename a rune or change its character there).  '
            'Yellow: not written yet.').font = f_sub

# --- write the master rows ----------------------------------------------------------------------
for i, row in enumerate(rows, start=2):
    rid, aff, p, rune, ch, oid = row[:6]
    for col, v in enumerate([rid, aff, p, rune], start=1):
        master.cell(row=i, column=col, value=v).font = f_input
    c = master.cell(row=i, column=5, value=ch)
    c.font, c.alignment = char_font(ch), center
    master.cell(row=i, column=6, value=oid).font = f_input
    master.cell(row=i, column=7, value=f'=IFERROR(INDEX($D:$D,MATCH($F{i},$A:$A,0)),"")').font = f_calc
    c = master.cell(row=i, column=8, value=f'=IFERROR(INDEX($E:$E,MATCH($F{i},$A:$A,0)),"")')
    c.font, c.alignment = char_font(ch, '000000'), center
    master.cell(row=i, column=9, value=skill_cell[rid]).font = f_link
    master.cell(row=i, column=10, value=f'U+{ord(ch):04X}').font = f_body
    for col in range(1, 11):
        master.cell(row=i, column=col).border = box
        if col not in (5, 8):
            master.cell(row=i, column=col).alignment = Alignment(vertical='center', wrap_text=(col == 9))
    master.row_dimensions[i].height = 24
master.auto_filter.ref = f'A1:J{len(rows) + 1}'
master['J1'].comment = Comment('Static: the code point the character had when the workbook was made. '
                               'It does not change if you type a new character.', 'Claude')
master['E1'].comment = Comment('Placeholder glyphs: the Runic block (Segoe UI Historic) for the five affinities '
                               'in the game, Alchemical Symbols (Segoe UI Symbol) for the rest. Type any character '
                               'to replace one; the affinity sheets follow.', 'Claude')

# --- Universal Runes -------------------------------------------------------------------------
UCOLS = ['ID', 'Rune', 'Character', 'Opposite ID', 'Opposite', 'Opposite character', 'As a means']
header(uni, 1, UCOLS, [10, 14, 11, 11, 14, 11, 60])
uni.freeze_panes = 'B2'
r = 2
k = 0
for n, ((ra, ma), (rb, mb)) in enumerate(universal, start=1):
    for side, rune, mean in (('a', ra, ma), ('b', rb, mb)):
        rid, oid = f'UNI-{n:02d}{side}', f'UNI-{n:02d}{"b" if side == "a" else "a"}'
        ch = UNI_CHARS[k]; k += 1
        uni.cell(row=r, column=1, value=rid).font = f_input
        uni.cell(row=r, column=2, value=rune).font = f_input
        c = uni.cell(row=r, column=3, value=ch)
        c.font, c.alignment = char_font(ch), center
        uni.cell(row=r, column=4, value=oid).font = f_input
        uni.cell(row=r, column=5, value=f'=IFERROR(INDEX($B:$B,MATCH($D{r},$A:$A,0)),"")').font = f_calc
        c = uni.cell(row=r, column=6, value=f'=IFERROR(INDEX($C:$C,MATCH($D{r},$A:$A,0)),"")')
        c.font, c.alignment = char_font(ch, '000000'), center
        uni.cell(row=r, column=7, value=mean).font = f_input
        for col in range(1, 8):
            uni.cell(row=r, column=col).border = box
            if col not in (3, 6):
                uni.cell(row=r, column=col).alignment = Alignment(vertical='center', wrap_text=(col == 7))
        uni.row_dimensions[r].height = 24
        r += 1
uni.auto_filter.ref = f'A1:G{r - 1}'

# --- Affinities index --------------------------------------------------------------------------
header(index, 1, ['Affinity', 'Rule', 'Status', 'Source', 'Base runes', 'Pairs', 'Sheet'], [14, 70, 24, 40, 11, 8, 12])
index.freeze_panes = 'B2'
r = 2
for aff, code, colour in ORDER:
    a = affinities[aff]
    index.cell(row=r, column=1, value=aff).font = Font(name=F, size=10, bold=True)
    index.cell(row=r, column=2, value=a['rule']).font = f_input
    index.cell(row=r, column=3, value=a['status']).font = f_input
    index.cell(row=r, column=4, value=a['source']).font = f_input
    index.cell(row=r, column=5, value=f"=COUNTIF('Base Runes'!$B:$B,$A{r})").font = f_link
    index.cell(row=r, column=6, value=f'=E{r}/2').font = f_calc
    c = index.cell(row=r, column=7, value=aff)
    c.hyperlink = f"#'{aff}'!A1"
    c.font = Font(name=F, size=10, color='0563C1', underline='single')
    for col in range(1, 8):
        index.cell(row=r, column=col).alignment = wrap
        index.cell(row=r, column=col).border = box
    r += 1
index.cell(row=r, column=1, value='Total').font = Font(name=F, size=10, bold=True)
index.cell(row=r, column=5, value=f'=SUM(E2:E{r - 1})').font = Font(name=F, size=10, bold=True)
index.cell(row=r, column=6, value=f'=SUM(F2:F{r - 1})').font = Font(name=F, size=10, bold=True)
r += 2
index.cell(row=r, column=1, value='Proposed, awaiting your yes or no (no sheets yet)').font = Font(name=F, size=10, bold=True)
r += 1
for p in PROPOSED:
    index.cell(row=r, column=1, value=p).font = f_input
    index.cell(row=r, column=3, value='Proposed').font = f_input
    r += 1
dv = DataValidation(type='list', formula1='"Deep dive v0,First pass,First pass (new 2026-10-22),Proposed,Approved,Cut"',
                    allow_blank=True)
index.add_data_validation(dv)
dv.add(f'C2:C{r}')

# --- Read Me -----------------------------------------------------------------------------------
lines = [
    ('Passage rune workbook', f_title),
    ('Made 2026-10-22 from design/14-rune-sets-v0.md and design/affinities/water.md. Nothing here is in the game yet.', f_sub),
    ('', f_body),
    ('Sheets', Font(name=F, size=11, bold=True)),
    ('Affinities: every affinity, its rule, status and rune count, with a link to its sheet. The proposed seven are listed below it.', f_body),
    ('Base Runes: every base rune in one list, with its character. This is where a rune\'s name and character live.', f_body),
    ('Universal Runes: the 28 universal pairs anyone can add to a skill, with characters.', f_body),
    ('One sheet per affinity: each rune as a skill, as a means, what it releases with its opposite, and the skills built on it.', f_body),
    ('', f_body),
    ('Colours', Font(name=F, size=11, bold=True)),
    ('Blue text: type here.', f_input),
    ('Green text: pulled from another sheet. Change it at its source (names and characters on Base Runes).', f_link),
    ('Black text: worked out on this sheet (opposites, counts).', f_calc),
    ('Yellow fill: not written yet (water is the only affinity with its means and together columns filled).', f_body),
    ('', f_body),
    ('Editing', Font(name=F, size=11, bold=True)),
    ('Rename a rune or change its character on Base Runes; its affinity sheet and opposite follow.', f_body),
    ('Add a rune: add a row on Base Runes with a new ID (for example FIR-6a, with its opposite FIR-6b), then a row on the affinity sheet with the same ID.', f_body),
    ('IDs: three-letter affinity code, pair number, a or b. The two runes of a pair point at each other in Opposite ID.', f_body),
    ('Characters are placeholders: the Runic block (font Segoe UI Historic) for the five affinities already in the game, Alchemical Symbols (Segoe UI Symbol) for the rest, and assorted symbols for universal runes. Type any character over them. In the game the ethercosm will draw original glyphs.', f_body),
    ('', f_body),
    ('Things to fix', Font(name=F, size=11, bold=True)),
    ('Rune names used by more than one affinity: ' + (', '.join(dupes) if dupes else 'none') + '. A rune name should be unique before this reaches the game.', f_body),
]
for i, (t, f) in enumerate(lines, start=1):
    c = readme.cell(row=i, column=1, value=t)
    c.font, c.alignment = f, Alignment(wrap_text=True, vertical='top')
readme.column_dimensions['A'].width = 120

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved', OUT, 'base runes', len(rows), 'universal', len(universal) * 2, 'dupes', dupes)
